"""Read-only audit of the author-designated workbook against manuscript Tables 2/3.

The workbook contains rounded seed-level metrics, not full training trajectories.
Do not treat numerical agreement as certification of historical executable code.
"""
import argparse
import hashlib
import json
import re
import statistics
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT.parents[1] / 'cifar10_3seed_sweep.xlsx'
OUT = ROOT / 'r4_tasl_experiments_20260921/legacy_main_sources_20261002'
ALGOS = ['FedAvg', 'Multi-Krum', 'Trimmed Mean', 'FLTrust', 'TASL']

def manuscript_rows(tex, label):
    block = tex.split('\\label{' + label + '}', 1)[1].split('\\end{table*}', 1)[0]
    return [line for line in block.splitlines()
            if any(line.startswith(a + ' ') for a in ALGOS)]

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--workbook', type=Path, default=DEFAULT_WORKBOOK)
    args = parser.parse_args()
    workbook = args.workbook
    w = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    seed_records = []
    for i, row in enumerate(w['All Seeds Data'].iter_rows(values_only=True), 1):
        if row[0] not in ALGOS or not str(row[2]).startswith('Seed '):
            continue
        seed_records.append({
            'algorithm': row[0], 'ratio_percent': int(str(row[1]).rstrip('%')),
            'seed': int(str(row[2]).split()[1]),
            'best': float(row[3]), 'avg': float(row[4]), 'asr': float(row[5]),
            'source_cells': {metric: f"'All Seeds Data'!{col}{i}"
                             for metric, col in [('best', 'D'), ('avg', 'E'), ('asr', 'F')]},
        })
    assert len(seed_records) == 60
    assert len({(r['algorithm'], r['ratio_percent'], r['seed']) for r in seed_records}) == 60
    tex = (ROOT / 'manuscript/main.tex').read_text(encoding='utf-8')
    indexed = []
    checks = []
    summaries = {}
    for metric, first_row in [('best', 4), ('avg', 14), ('asr', 24)]:
        for ai, algo in enumerate(ALGOS):
            for ri, ratio in enumerate([10, 20, 30, 40]):
                series = [r for r in seed_records
                          if r['algorithm'] == algo and r['ratio_percent'] == ratio]
                assert sorted(r['seed'] for r in series) == [42, 123, 2024]
                values = [r[metric] for r in series]
                mean, sd = statistics.mean(values), statistics.stdev(values)
                cells = [f'{get_column_letter(2+ri*2+j)}{first_row+ai}' for j in [0, 1]]
                report = [float(w['Summary Matrix'][c].value) for c in cells]
                # Both source per-seed values and summary are printed to 0.01 point.
                # Mean error <= .01; SD has a slightly wider rounding allowance.
                assert abs(mean-report[0]) < .011, (algo, ratio, metric, 'mean')
                assert abs(sd-report[1]) < .016, (algo, ratio, metric, 'sd')
                summaries[(algo, ratio, metric)] = report
                checks.append({'algorithm': algo, 'ratio_percent': ratio, 'metric': metric,
                    'seed_values': values, 'source_seed_cells': [r['source_cells'][metric] for r in series],
                    'recomputed_mean': mean, 'recomputed_sample_sd': sd,
                    'workbook_mean': report[0], 'workbook_sample_sd': report[1],
                    'source_summary_cells': ["'Summary Matrix'!"+c for c in cells]})
    table2 = manuscript_rows(tex, 'tab:robustness_combined')
    assert len(table2) == 5
    for ai, algo in enumerate(ALGOS):
        part = table2[ai].split('&')[4]
        submitted = [float(v) for v in re.findall(r'\d+(?:\.\d+)?', part)]
        assert len(submitted) == 2
        record = next(r for r in seed_records if r['algorithm'] == algo
                      and r['ratio_percent'] == 40 and r['seed'] == 42)
        for metric, value in zip(['best', 'asr'], submitted):
            assert abs(value-record[metric]) <= .055, (algo, metric, value, record[metric])
            indexed.append({'table': 2, 'dataset': 'cifar10', 'attack': 'label_flip',
                'algorithm': algo, 'ratio_percent': 40, 'metric': metric, 'statistic': 'seed42',
                'manuscript_value': value, 'workbook_value': record[metric],
                'source_cell': record['source_cells'][metric], 'source_precision': '0.01 percentage point'})
    table3 = manuscript_rows(tex, 'tab:ratio_sweep')
    assert len(table3) == 15
    for mi, metric in enumerate(['best', 'avg', 'asr']):
        for ai, algo in enumerate(ALGOS):
            parts = table3[mi*5+ai].split('&')[5:9]
            assert len(parts) == 4
            for ri, (ratio, part) in enumerate(zip([10, 20, 30, 40], parts)):
                submitted = [float(v) for v in re.findall(r'\d+(?:\.\d+)?', part)]
                assert len(submitted) == 2
                source = summaries[(algo, ratio, metric)]
                check = next(c for c in checks if c['algorithm'] == algo
                             and c['ratio_percent'] == ratio and c['metric'] == metric)
                for j, statistic in enumerate(['mean', 'sample_sd']):
                    assert abs(submitted[j]-source[j]) <= .055, (algo, ratio, metric, statistic)
                    indexed.append({'table': 3, 'dataset': 'cifar10', 'attack': 'label_flip',
                        'algorithm': algo, 'ratio_percent': ratio, 'metric': metric,
                        'statistic': statistic, 'manuscript_value': submitted[j],
                        'workbook_value': source[j], 'source_cell': check['source_summary_cells'][j],
                        'source_precision': '0.01 percentage point'})
    assert len(indexed) == 130
    result = {'validation': 'passed', 'comments': ['R4.10', 'R4.15'],
        'workbook_filename': workbook.name,
        'workbook_sha256': hashlib.sha256(workbook.read_bytes()).hexdigest(),
        'seed_level_records': seed_records, 'three_seed_checks': checks, 'cell_index': indexed,
        'scope': {'Table2_CIFAR_LF_numeric_values': 10, 'Table3_CIFAR_numeric_values': 120},
        'limitations': ['Workbook seed-level values and summaries are rounded records, not unrounded trajectories.',
            'Historical code versions, partitions, and execution environment are not certified by this check.',
            'MNIST/Fashion-MNIST main tables and CIFAR-10 SF/Scaling main-table trajectories remain unrecovered.']}
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / 'workbook_validation.json').write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding='utf-8')
    print(json.dumps({'validation': 'passed', 'seed_records': 60, 'three_seed_metric_checks': 60,
                      'indexed_numeric_values': len(indexed), 'workbook_modified': False}))

if __name__ == '__main__':
    main()
